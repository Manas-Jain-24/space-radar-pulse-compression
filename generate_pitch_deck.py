"""Generate the 5-Slide Pitch Deck (.pptx) for Utkrishta 2026 Track 1 using python-pptx."""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck(filename="PITCH_DECK.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 Widescreen
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette (Space-Tech Navy, Cyan, Gold, White)
    DARK_BG = RGBColor(10, 18, 36)
    CARD_BG = RGBColor(18, 28, 52)
    CYAN_ACCENT = RGBColor(0, 212, 255)
    GOLD_ACCENT = RGBColor(255, 193, 7)
    WHITE = RGBColor(255, 255, 255)
    GRAY = RGBColor(180, 195, 215)
    BORDER_COLOR = RGBColor(35, 55, 95)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="TRACK 1: RADAR SIGNAL PROCESSING | SDSC SHAR"):
        # Category pill
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = CYAN_ACCENT

        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.7))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE

    # =========================================================================
    # SLIDE 1: TITLE & TEAM
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Decorative banner box
    decor = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.0), Inches(10.933), Inches(5.5))
    decor.fill.solid()
    decor.fill.fore_color.rgb = CARD_BG
    decor.line.color.rgb = BORDER_COLOR
    decor.line.width = Pt(1.5)

    tb1 = s1.shapes.add_textbox(Inches(1.6), Inches(1.4), Inches(10.133), Inches(4.7))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "UTKRISHTA 2026 | SSG × RESEARCH CLUB, IIIT SRI CITY"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = CYAN_ACCENT

    p1 = tf1.add_paragraph()
    p1.text = "LFM Radar Pulse Compression & +10 dB SNR Boost for Long-Range Space Target Tracking"
    p1.font.size = Pt(26)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_before = Pt(14)
    p1.space_after = Pt(10)

    p2 = tf1.add_paragraph()
    p2.text = "Solving the Radar Energy-Resolution Dilemma at 200 km Range with Linear Frequency Modulation"
    p2.font.size = Pt(14)
    p2.font.color.rgb = GOLD_ACCENT
    p2.space_after = Pt(24)

    p3 = tf1.add_paragraph()
    p3.text = "TRACK 1: Radar Signal Processing — Target Detection, Pulse Compression & SNR Enhancement\nCurated by: Mr. Rajendra Kumar, Scientist, Radar Division, SDSC SHAR\n\nTeam Name: Team Beyond Signals | Event: 24-Hour Research-Athon (Utkrishta 2026)"
    p3.font.size = Pt(12)
    p3.font.color.rgb = GRAY

    # =========================================================================
    # SLIDE 2: PROBLEM OVERVIEW & RADAR SETUP
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Problem Overview & Space-Tracking System Setup")

    # Left Card: The Radar Dilemma
    c1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = BORDER_COLOR
    
    tb2a = s2.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
    tf2a = tb2a.text_frame
    tf2a.word_wrap = True
    p = tf2a.paragraphs[0]
    p.text = "The Long-Range Radar Dilemma"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    
    points_a = [
        "Mission Requirement: Detect distant, small RCS space objects (R = 200 km, RCS = 1.0 m²) at SDSC SHAR.",
        "Classical CW Pulse Conflict: Higher detection range requires high pulse energy Et = Pt · tau. But simple pulses tie range resolution directly to pulse duration: dR = c · tau / 2.",
        "Catastrophic Resolution Loss: A 100 μs unmodulated pulse yields dR = 15 km (unusable for precision tracking).",
        "The LFM Solution: Linearly chirp the frequency across bandwidth B = 4 MHz. Resolution becomes dR = c / (2B) = 37.5 m, independent of pulse duration!"
    ]
    for pt in points_a:
        p = tf2a.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = WHITE
        p.space_before = Pt(8)

    # Right Card: System Parameters
    c2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BORDER_COLOR

    tb2b = s2.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.8))
    tf2b = tb2b.text_frame
    tf2b.word_wrap = True
    p = tf2b.paragraphs[0]
    p.text = "SDSC SHAR Radar Link Specifications"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = GOLD_ACCENT

    specs = [
        ("Carrier Frequency (fc)", "3.0 GHz (S-Band, λ = 0.1 m)"),
        ("Pulse Repetition Time (PRT)", "2.0 ms (PRF = 500 Hz, Runamb = 300 km)"),
        ("Chirp Sweep Bandwidth (B)", "4.0 MHz (Nominal dR = 37.5 m)"),
        ("Pulse Widths Compared", "10 μs (TB = 40) vs 100 μs (TB = 400)"),
        ("Transmit Peak Power (Pt)", "250 kW (53.98 dBW)"),
        ("Antenna Gain (Tx & Rx)", "40.0 dBi each (Beamwidth ~ 1.8°)"),
        ("Target Range & RCS", "R = 200.0 km, RCS = 1.0 m² (0 dBsm)"),
        ("Received Power (Pr)", "3.95 × 10⁻¹⁴ W (-134.04 dBW)"),
        ("Noise Figure & Losses", "F = 2.5 dB, System Loss L = 3.0 dB")
    ]
    for k, v in specs:
        p = tf2b.add_paragraph()
        p.text = f"• {k}: {v}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = WHITE
        p.space_before = Pt(4)

    # =========================================================================
    # SLIDE 3: METHODOLOGY & SIMULATION WORKFLOW
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Simulation Methodology & Mathematical Pipeline")

    # 4 Workflow Step Boxes
    steps = [
        ("Step 1: Signal Synthesis", "PRT = 2 ms, B = 4 MHz, fs = 16 MHz. Dual LFM up-chirps: s(t) = rect(t/tau) exp(j π kr t²). Slope kr1 = 4e11 Hz/s vs kr2 = 4e10 Hz/s."),
        ("Step 2: Channel & Propagation", "Radar Range Equation at R = 200 km (delay 1.333 ms). Circular complex AWGN with discrete-time variance Pn = N0 · fs."),
        ("Step 3: Matched Filter (MF)", "Zero-padded FFT linear cross-correlation: y(t) = ifft(fft(rx) · conj(fft(ref))). Trims circular convolution artifacts."),
        ("Step 4: Detection & CA-CFAR", "1-D Cell-Averaging CFAR (Nref=48, Nguard=16, Pfa=1e-6). Link margin verification against Swerling 0 Marcum Q model.")
    ]
    
    for idx, (st_title, st_desc) in enumerate(steps):
        x = Inches(0.8 + idx * 2.95)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.6), Inches(2.8), Inches(2.5))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CYAN_ACCENT
        card.line.width = Pt(1.0)
        
        tb = s3.shapes.add_textbox(x + Inches(0.1), Inches(1.7), Inches(2.6), Inches(2.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = st_title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = GOLD_ACCENT
        
        p2 = tf.add_paragraph()
        p2.text = st_desc
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = WHITE
        p2.space_before = Pt(6)

    # Embed Figure 2 in lower half
    fig2_path = Path("figures/track1/fig2_lfm_waveforms_and_spectra.png")
    if fig2_path.exists():
        s3.shapes.add_picture(str(fig2_path), Inches(1.8), Inches(4.3), width=Inches(9.7))

    # =========================================================================
    # SLIDE 4: VISUAL PLOTS & BENCHMARK TABLE
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Quantitative Verification & Experimental Plots")

    # Left: Embed Fig 4 (Pulse Compression Zoom)
    fig4_path = Path("figures/track1/fig4_pulse_compression_snr_boost.png")
    if fig4_path.exists():
        s4.shapes.add_picture(str(fig4_path), Inches(0.8), Inches(1.6), width=Inches(6.2))

    # Right: Benchmark Table
    rows = [
        ("Parameter / Metric", "Pulse 1 (10 μs)", "Pulse 2 (100 μs)", "Gain / Ratio"),
        ("Time-Bandwidth (TB)", "40 (16.0 dB)", "400 (26.0 dB)", "+10.0 dB"),
        ("Radiated Energy", "2.50 J", "25.00 J", "+10.0 dB"),
        ("Theoretical MF SNR", "17.44 dB", "27.44 dB", "+10.00 dB"),
        ("Simulated Peak SNR", "17.13 dB", "27.31 dB", "+10.18 dB (Valid!)"),
        ("Nominal Resolution", "37.50 m", "37.50 m", "1.00x (Identical)"),
        ("Measured -3dB Width", "35.36 m", "34.17 m", "0.97x (Identical)"),
        ("Required SNR (Pd=0.9)", "13.18 dB", "13.18 dB", "Marcum Q baseline"),
        ("Link Margin", "+4.25 dB (Marginal)", "+14.25 dB (Robust)", "+10.00 dB Boost"),
        ("CA-CFAR Margin", "+6.81 dB", "+16.99 dB", "+10.18 dB Boost")
    ]
    
    t_shape = s4.shapes.add_table(len(rows), 4, Inches(7.2), Inches(1.6), Inches(5.4), Inches(5.2))
    table = t_shape.table
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(1.1)
    table.columns[2].width = Inches(1.1)
    table.columns[3].width = Inches(1.2)

    for r_idx, row_vals in enumerate(rows):
        for c_idx, val in enumerate(row_vals):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(8.5)
            if r_idx == 0:
                p.font.bold = True
                p.font.color.rgb = CYAN_ACCENT
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(25, 45, 80)
            else:
                p.font.color.rgb = WHITE if c_idx < 3 else GOLD_ACCENT
                cell.fill.solid()
                cell.fill.fore_color.rgb = CARD_BG

    # =========================================================================
    # SLIDE 5: INSIGHTS, BONUS EXPERIMENTS & CONCLUSION
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Operational Insights, Trade-offs & Conclusions")

    # Left: Insights & Trade-offs
    c5a = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    c5a.fill.solid()
    c5a.fill.fore_color.rgb = CARD_BG
    c5a.line.color.rgb = BORDER_COLOR

    tb5a = s5.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
    tf5a = tb5a.text_frame
    tf5a.word_wrap = True
    p = tf5a.paragraphs[0]
    p.text = "Operational Insights & Engineering Trade-offs"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    insights = [
        "Rigorous +10 dB Proof: We mathematically, visually, and empirically verified that expanding pulse duration by 10x yields a +10.18 dB SNR boost while strictly maintaining identical ~35 m range resolution.",
        "Operational Impact: At 200 km, the 10 μs pulse provides only +4.25 dB link margin (vulnerable to rain fade or pointing loss). The 100 μs pulse provides an ironclad +14.25 dB margin.",
        "Blind Zone Limitation (Rmin = c·tau / 2): A 100 μs pulse blinds the radar up to 15.0 km (vs 1.5 km for 10 μs). Operational tracking radars solve this with dual-waveform scheduling.",
        "Doppler Coupling: For high-speed space targets, radial velocity shifts the peak by dt = -fd / kr. Alternating up/down chirps resolves this ambiguity."
    ]
    for ins in insights:
        p = tf5a.add_paragraph()
        p.text = "• " + ins
        p.font.size = Pt(10)
        p.font.color.rgb = WHITE
        p.space_before = Pt(6)

    # Right: Embed Fig 5 (CA-CFAR)
    fig5_path = Path("figures/track1/fig5_detection_and_cfar_margin.png")
    if fig5_path.exists():
        s5.shapes.add_picture(str(fig5_path), Inches(6.8), Inches(1.6), width=Inches(5.7))

    prs.save(filename)
    print(f"Successfully generated 5-Slide Pitch Deck: {Path(filename).resolve()}")

if __name__ == "__main__":
    create_deck()
